"""
generar_plantillas.py
------------------------
Genera las plantillas Excel vacías (plantillas/plantilla_ventas.xlsx y
plantilla_compras.xlsx) para que ustedes o el cliente carguen los
comprobantes a mano cuando no hay export de AFIP disponible.

Se corre una sola vez (o cuando quieran regenerar las plantillas):
    python generar_plantillas.py
"""

import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side

FUENTE = "Arial"
HEADER_FONT = Font(name=FUENTE, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
EJEMPLO_FILL = PatternFill("solid", fgColor="FFF2CC")
LEYENDA_FONT = Font(name=FUENTE, italic=True, size=9, color="666666")
BORDE = Border(*(Side(style="thin", color="B7B7B7"),) * 4)

ENCABEZADOS = [
    "fecha", "tipo_comprobante", "punto_venta", "numero_comprobante",
    "cuit", "razon_social", "neto_gravado", "iva",
    "percepcion_iva", "percepcion_iibb",
]
EJEMPLO_VENTA = ["2026-08-05", "Factura A", 1, 145, "20304050607", "Cliente Ejemplo SA", 100000, 21000, 0, 0]
EJEMPLO_COMPRA = ["2026-08-03", "Factura A", 3, 8821, "30111222338", "Proveedor Ejemplo SRL", 50000, 10500, 0, 1500]


def _armar_plantilla(path: str, ejemplo: list):
    wb = Workbook()
    ws = wb.active
    ws.title = "Comprobantes"

    for col_idx, encabezado in enumerate(ENCABEZADOS, start=1):
        celda = ws.cell(row=1, column=col_idx, value=encabezado)
        celda.font = HEADER_FONT
        celda.fill = HEADER_FILL
        celda.border = BORDE

    for col_idx, valor in enumerate(ejemplo, start=1):
        celda = ws.cell(row=2, column=col_idx, value=valor)
        celda.fill = EJEMPLO_FILL
        celda.border = BORDE

    ws.cell(row=4, column=1,
             value="↑ La fila amarilla es un EJEMPLO: bórrenla y carguen sus propios comprobantes desde la fila 2.")
    ws.cell(row=4, column=1).font = LEYENDA_FONT
    ws.merge_cells(start_row=4, start_column=1, end_row=4, end_column=10)

    ws.cell(row=5, column=1,
             value="fecha: AAAA-MM-DD | neto_gravado e iva: números sin puntos de miles, con coma o punto decimal.")
    ws.cell(row=5, column=1).font = LEYENDA_FONT
    ws.merge_cells(start_row=5, start_column=1, end_row=5, end_column=10)

    ws.cell(row=6, column=1,
             value="punto_venta y numero_comprobante: tal cual figuran impresos en el comprobante "
                   "(ej. punto_venta 1, numero_comprobante 145). Son obligatorios si van a generar "
                   "los TXT para el Libro IVA Digital de AFIP; para la liquidación de IVA sola no "
                   "hacen falta.")
    ws.cell(row=6, column=1).font = LEYENDA_FONT
    ws.merge_cells(start_row=6, start_column=1, end_row=6, end_column=10)

    ws.cell(row=7, column=1,
             value="percepcion_iva y percepcion_iibb: OPCIONALES — solo si ese comprobante puntual "
                   "tiene una percepción. Si no tiene, dejar en 0 o en blanco.")
    ws.cell(row=7, column=1).font = LEYENDA_FONT
    ws.merge_cells(start_row=7, start_column=1, end_row=7, end_column=10)

    anchos = [14, 20, 14, 20, 16, 30, 16, 14, 16, 16]
    from openpyxl.utils import get_column_letter
    for col_idx, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = ancho

    wb.save(path)


def main():
    os.makedirs("plantillas", exist_ok=True)
    _armar_plantilla("plantillas/plantilla_ventas.xlsx", EJEMPLO_VENTA)
    _armar_plantilla("plantillas/plantilla_compras.xlsx", EJEMPLO_COMPRA)
    print("Listo: plantillas/plantilla_ventas.xlsx y plantillas/plantilla_compras.xlsx")


if __name__ == "__main__":
    main()
