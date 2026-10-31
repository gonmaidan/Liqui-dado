"""
reporte_excel.py
------------------
Genera el Excel de liquidación de IVA: una hoja "Resumen" con la
determinación del impuesto (con fórmulas, no valores fijos) y dos hojas
de detalle (Ventas / Compras) con los comprobantes cargados.
"""

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from calculo_iva import ResultadoIVA

FUENTE = "Arial"
AZUL_INPUT = Font(name=FUENTE, color="0000FF")
NEGRO = Font(name=FUENTE, color="000000")
NEGRO_BOLD = Font(name=FUENTE, bold=True, color="000000")
TITULO = Font(name=FUENTE, bold=True, size=14, color="FFFFFF")
FILL_TITULO = PatternFill("solid", fgColor="1F4E78")
FILL_HEADER = PatternFill("solid", fgColor="D9E1F2")
FILL_SALDO = PatternFill("solid", fgColor="FFF2CC")
BORDE = Border(*(Side(style="thin", color="B7B7B7"),) * 4)
FORMATO_MONEDA = '$#,##0.00;($#,##0.00);"-"'


def _hoja_detalle(wb: Workbook, nombre: str, df: pd.DataFrame):
    ws = wb.create_sheet(nombre)
    columnas = ["fecha", "tipo_comprobante", "cuit", "razon_social", "neto_gravado", "iva"]
    encabezados = ["Fecha", "Tipo Comprobante", "CUIT", "Razón Social", "Neto Gravado", "IVA"]

    for col_idx, encabezado in enumerate(encabezados, start=1):
        celda = ws.cell(row=1, column=col_idx, value=encabezado)
        celda.font = NEGRO_BOLD
        celda.fill = FILL_HEADER
        celda.border = BORDE

    for row_idx, (_, fila) in enumerate(df.iterrows(), start=2):
        for col_idx, col in enumerate(columnas, start=1):
            valor = fila[col]
            celda = ws.cell(row=row_idx, column=col_idx, value=valor)
            celda.font = NEGRO
            celda.border = BORDE
            if col in ("neto_gravado", "iva"):
                celda.number_format = FORMATO_MONEDA

    ultima_fila = len(df) + 1
    if ultima_fila >= 2:
        fila_total = ultima_fila + 1
        ws.cell(row=fila_total, column=4, value="Total").font = NEGRO_BOLD
        for col_idx, col in ((5, "neto_gravado"), (6, "iva")):
            letra = get_column_letter(col_idx)
            celda = ws.cell(row=fila_total, column=col_idx,
                             value=f"=SUM({letra}2:{letra}{ultima_fila})")
            celda.font = NEGRO_BOLD
            celda.number_format = FORMATO_MONEDA

    anchos = [14, 20, 16, 30, 16, 14]
    for col_idx, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = ancho

    return ultima_fila  # última fila de datos (sin contar el total), para referenciar desde Resumen


def generar_excel(resultado: ResultadoIVA, ventas: pd.DataFrame, compras: pd.DataFrame, path_salida: str):
    wb = Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("Resumen", 0)
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 20

    ws.merge_cells("A1:B1")
    ws["A1"] = f"Liquidación de IVA — Período {resultado.periodo}"
    ws["A1"].font = TITULO
    ws["A1"].fill = FILL_TITULO
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 26

    # Genero primero las hojas de detalle para saber hasta qué fila referenciar
    ultima_fila_ventas = _hoja_detalle(wb, "Ventas", ventas)
    ultima_fila_compras = _hoja_detalle(wb, "Compras", compras)

    filas = [
        ("Total Ventas (Neto Gravado)", f"=SUM(Ventas!E2:E{ultima_fila_ventas})", False),
        ("Débito Fiscal (IVA Ventas)", f"=SUM(Ventas!F2:F{ultima_fila_ventas})", False),
        ("Total Compras (Neto Gravado)", f"=SUM(Compras!E2:E{ultima_fila_compras})", False),
        ("Crédito Fiscal (IVA Compras)", f"=SUM(Compras!F2:F{ultima_fila_compras})", False),
    ]

    fila_actual = 3
    for etiqueta, formula, _ in filas:
        ws.cell(row=fila_actual, column=1, value=etiqueta).font = NEGRO
        celda_valor = ws.cell(row=fila_actual, column=2, value=formula)
        celda_valor.font = NEGRO
        celda_valor.number_format = FORMATO_MONEDA
        fila_actual += 1

    fila_saldo = fila_actual + 1
    ws.cell(row=fila_saldo, column=1, value="Saldo de IVA (a pagar / a favor)").font = NEGRO_BOLD
    celda_saldo = ws.cell(row=fila_saldo, column=2, value="=B4-B6")  # Débito - Crédito
    celda_saldo.font = NEGRO_BOLD
    celda_saldo.number_format = FORMATO_MONEDA
    celda_saldo.fill = FILL_SALDO
    ws.cell(row=fila_saldo, column=1).fill = FILL_SALDO

    nota = ws.cell(row=fila_saldo + 2, column=1,
                    value="Saldo positivo = IVA a pagar. Saldo negativo = saldo técnico a favor del período.")
    nota.font = Font(name=FUENTE, italic=True, size=9, color="666666")
    ws.merge_cells(start_row=fila_saldo + 2, start_column=1, end_row=fila_saldo + 2, end_column=4)

    wb.save(path_salida)
