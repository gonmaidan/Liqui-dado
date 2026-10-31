"""
cm.py
-------
Liquidación mensual de Convenio Multilateral (CM03) por cliente.

  /clientes/<id>/cm                    -> estado de la configuración y elegir período
  /clientes/<id>/cm/config             -> jurisdicciones (+ mínimo), actividades y alícuotas
  /clientes/<id>/cm/coeficientes       -> coeficientes unificados (CM05) por vigencia
  /clientes/<id>/cm/<AAAA-MM>          -> ingresos, deducciones y resultado
  /clientes/<id>/cm/<AAAA-MM>/pdf|xlsx -> descargas

Las alícuotas, mínimos y coeficientes los carga el estudio (no se inventan). Los ingresos del
mes se proponen desde las facturas de venta cargadas y las percepciones de IIBB desde las de compra.
"""

import io
import os
import re
import uuid

from flask import Blueprint, abort, flash, redirect, render_template, request, send_file, url_for
from flask_login import current_user

from auth import estudio_required
from clientes import _cliente_del_estudio_o_404
from cm_calculo import (
    JURISDICCIONES, calcular_cm03, canonizar_jurisdiccion, coeficientes_vigentes, validar_coeficientes,
)
from extensions import db
from f931 import _num
from models import (
    CmActividad, CmAlicuota, CmCoeficiente, CmCredito, CmIngreso, CmJurisdiccion, Comprobante,
)

cm_bp = Blueprint("cm", __name__, url_prefix="/clientes/<int:cliente_id>/cm")

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
TMP_DIR = os.path.join(BASE_DIR, "_tmp_web")
os.makedirs(TMP_DIR, exist_ok=True)
RE_PERIODO = re.compile(r"^\d{4}-(0[1-9]|1[0-2])$")
CAMPOS_CREDITO = ("retenciones", "sircreb", "percepciones", "pagos", "saldo_anterior")


# ------------------------------------------------------------------ datos del cliente
def _periodo_o_404(periodo):
    if not RE_PERIODO.match(periodo):
        abort(404)
    return periodo


def _jurisdicciones_cliente(cliente):
    orden = {j: i for i, j in enumerate(JURISDICCIONES)}
    return sorted(cliente.cm_jurisdicciones, key=lambda x: orden.get(x.jurisdiccion, 99))


def _conjuntos_coeficientes(cliente):
    por_vigencia = {}
    for c in cliente.cm_coeficientes:
        por_vigencia.setdefault(c.vigencia_desde, {})[c.jurisdiccion] = c.coeficiente
    return sorted(por_vigencia.items())


def _es_nota_de_credito(c) -> bool:
    t = (c.tipo_comprobante or "").lower()
    return "crédito" in t or "credito" in t


def _ventas_del_periodo(cliente, periodo):
    """(neto de ventas del mes con notas de crédito restadas, cantidad de comprobantes, cantidad de NC)."""
    total, n, nc = 0.0, 0, 0
    for c in Comprobante.query.filter_by(cliente_id=cliente.id, tipo="venta", periodo=periodo):
        signo = -1 if _es_nota_de_credito(c) else 1
        total += signo * float(c.neto_gravado)
        n += 1
        nc += 1 if signo < 0 else 0
    return round(total, 2), n, nc


def _percepciones_de_compras(cliente, periodo, habilitadas):
    """({jurisdicción habilitada: importe}, sin_identificar, {jurisdicción no habilitada: importe})."""
    por_jur, sin_id, ajenas = {}, 0.0, {}
    for c in Comprobante.query.filter_by(cliente_id=cliente.id, tipo="compra", periodo=periodo):
        signo = -1 if _es_nota_de_credito(c) else 1
        filas = [(p.jurisdiccion, float(p.importe)) for p in c.percepciones_iibb]
        if not filas and float(c.percepcion_iibb or 0):
            filas = [(c.jurisdiccion_iibb, float(c.percepcion_iibb))]
        for nombre, importe in filas:
            jur = canonizar_jurisdiccion(nombre)
            if jur is None:
                sin_id += signo * importe
            elif jur in habilitadas:
                por_jur[jur] = por_jur.get(jur, 0.0) + signo * importe
            else:
                ajenas[jur] = ajenas.get(jur, 0.0) + signo * importe
    return ({k: round(v, 2) for k, v in por_jur.items()}, round(sin_id, 2), {k: round(v, 2) for k, v in ajenas.items()})


def _contexto_periodo(cliente, periodo):
    """Todo lo necesario para mostrar/calcular el CM03 del período."""
    jurs = _jurisdicciones_cliente(cliente)
    habilitadas = [j.jurisdiccion for j in jurs]
    actividades = list(cliente.cm_actividades)
    ventas, n_ventas, n_nc = _ventas_del_periodo(cliente, periodo)
    percep, sin_id, ajenas = _percepciones_de_compras(cliente, periodo, habilitadas)

    ingresos_guardados = {i.actividad_id: float(i.importe) for i in cliente.cm_ingresos if i.periodo == periodo}
    hay_ingresos_guardados = bool(ingresos_guardados)
    ingresos = {}
    for k, a in enumerate(actividades):
        if hay_ingresos_guardados:
            ingresos[a.id] = ingresos_guardados.get(a.id, 0.0)
        else:
            ingresos[a.id] = ventas if k == 0 else 0.0

    creditos_guardados = {c.jurisdiccion: c for c in cliente.cm_creditos if c.periodo == periodo}
    creditos = {}
    for j in habilitadas:
        g = creditos_guardados.get(j)
        if g:
            creditos[j] = {k: float(getattr(g, k)) for k in CAMPOS_CREDITO}
        else:
            creditos[j] = {k: 0.0 for k in CAMPOS_CREDITO} | {"percepciones": percep.get(j, 0.0)}
    return dict(
        jurs=jurs, habilitadas=habilitadas, actividades=actividades, ventas=ventas, n_ventas=n_ventas, n_nc=n_nc,
        percep=percep, percep_sin_identificar=sin_id, percep_ajenas=ajenas, ingresos=ingresos,
        ingresos_guardados=hay_ingresos_guardados, creditos=creditos,
        creditos_guardados=bool(creditos_guardados),
    )


def _faltantes_de_configuracion(cliente, periodo):
    """Lista de cosas que faltan configurar para poder calcular (vacía si está todo)."""
    faltan = []
    if not cliente.cm_jurisdicciones:
        faltan.append("Elegí las jurisdicciones donde tributa el cliente.")
    if not cliente.cm_actividades:
        faltan.append("Cargá al menos una actividad con sus alícuotas.")
    vig, coefs = coeficientes_vigentes(_conjuntos_coeficientes(cliente), periodo)
    if coefs is None:
        faltan.append(f"No hay coeficientes unificados vigentes para {periodo}. Cargalos en «Coeficientes».")
    return faltan


def _calcular(cliente, periodo, ctx=None):
    ctx = ctx or _contexto_periodo(cliente, periodo)
    vig, coefs = coeficientes_vigentes(_conjuntos_coeficientes(cliente), periodo)
    ingresos = {a.nombre: ctx["ingresos"][a.id] for a in ctx["actividades"]}
    alicuotas = {(al.actividad.nombre, al.jurisdiccion): al.alicuota for a in ctx["actividades"] for al in a.alicuotas}
    minimos = {j.jurisdiccion: float(j.minimo) for j in ctx["jurs"]}
    resultado = calcular_cm03(periodo, ingresos, coefs, alicuotas, minimos, ctx["creditos"], vigencia_coef=vig)
    for j in coefs:
        if j not in ctx["habilitadas"]:
            resultado.avisos.append(f"El conjunto de coeficientes incluye {j}, que no está entre las jurisdicciones del cliente.")
    for j in ctx["habilitadas"]:
        if coefs.get(j) is None:
            resultado.avisos.append(f"Falta el coeficiente de {j} en el conjunto vigente ({vig}).")
    return resultado


# ------------------------------------------------------------------ pantallas
@cm_bp.route("")
@estudio_required
def inicio(cliente_id):
    cliente = _cliente_del_estudio_o_404(cliente_id)
    conjuntos = _conjuntos_coeficientes(cliente)
    from datetime import date
    hoy = date.today()
    ant = f"{hoy.year - 1}-12" if hoy.month == 1 else f"{hoy.year}-{hoy.month - 1:02d}"
    return render_template(
        "cm_inicio.html", cliente=cliente, jurs=_jurisdicciones_cliente(cliente), actividades=cliente.cm_actividades,
        conjuntos=conjuntos, periodo_sugerido=request.args.get("periodo") or ant,
    )


@cm_bp.route("/ir", methods=["POST"])
@estudio_required
def ir(cliente_id):
    cliente = _cliente_del_estudio_o_404(cliente_id)
    periodo = request.form.get("periodo", "").strip()
    if not RE_PERIODO.match(periodo):
        flash("El período tiene que tener el formato AAAA-MM (por ejemplo 2026-09).")
        return redirect(url_for("cm.inicio", cliente_id=cliente.id))
    return redirect(url_for("cm.liquidacion", cliente_id=cliente.id, periodo=periodo))


@cm_bp.route("/config", methods=["GET", "POST"])
@estudio_required
def config(cliente_id):
    cliente = _cliente_del_estudio_o_404(cliente_id)
    if request.method == "POST":
        f = request.form
        try:
            # --- jurisdicciones y mínimos
            elegidas = [j for j in JURISDICCIONES if f.get("jur_" + j) == "1"]
            if not elegidas:
                raise ValueError("Tildá al menos una jurisdicción.")
            minimos = {j: _num(f.get("min_" + j)) for j in elegidas}
            if any(v < 0 for v in minimos.values()):
                raise ValueError("Los mínimos no pueden ser negativos.")
            existentes = {x.jurisdiccion: x for x in cliente.cm_jurisdicciones}
            for j in list(existentes):
                if j not in elegidas:
                    db.session.delete(existentes[j])
                    for act in cliente.cm_actividades:
                        for al in list(act.alicuotas):
                            if al.jurisdiccion == j:
                                db.session.delete(al)
            for j in elegidas:
                if j in existentes:
                    existentes[j].minimo = minimos[j]
                else:
                    db.session.add(CmJurisdiccion(cliente_id=cliente.id, jurisdiccion=j, minimo=minimos[j]))

            # --- actividades existentes
            nombres_usados = set()
            for act in list(cliente.cm_actividades):
                if f.get(f"borrar_{act.id}") == "1":
                    db.session.delete(act)
                    continue
                nombre = f.get(f"nombre_{act.id}", "").strip()[:80]
                if not nombre:
                    raise ValueError("Una actividad quedó sin nombre (para quitarla, tildá «Borrar»).")
                if nombre.lower() in nombres_usados:
                    raise ValueError(f"La actividad «{nombre}» está repetida.")
                nombres_usados.add(nombre.lower())
                act.nombre = nombre
                _guardar_alicuotas(act, elegidas, lambda j, a=act: f.get(f"alic_{a.id}_{j}"))

            # --- actividad nueva
            nueva = f.get("nueva_nombre", "").strip()[:80]
            if nueva:
                if nueva.lower() in nombres_usados:
                    raise ValueError(f"La actividad «{nueva}» está repetida.")
                act = CmActividad(cliente_id=cliente.id, nombre=nueva, orden=len(cliente.cm_actividades) + 1)
                db.session.add(act)
                db.session.flush()
                _guardar_alicuotas(act, elegidas, lambda j: f.get(f"alic_nueva_{j}"))
        except ValueError as e:
            db.session.rollback()
            flash(str(e) if len(str(e)) > 12 else "Alguno de los importes o alícuotas no es un número válido.")
            return redirect(url_for("cm.config", cliente_id=cliente.id))
        db.session.commit()
        flash("Configuración guardada.")
        return redirect(url_for("cm.config", cliente_id=cliente.id))

    alic = {(al.actividad_id, al.jurisdiccion): al.alicuota for a in cliente.cm_actividades for al in a.alicuotas}
    return render_template(
        "cm_config.html", cliente=cliente, todas=JURISDICCIONES,
        elegidas={j.jurisdiccion: float(j.minimo) for j in cliente.cm_jurisdicciones},
        jurs=_jurisdicciones_cliente(cliente), actividades=cliente.cm_actividades, alic=alic,
    )


def _guardar_alicuotas(actividad, jurisdicciones, leer):
    """Crea/actualiza/borra las alícuotas de la actividad según lo escrito (vacío = no tiene cargada)."""
    actuales = {al.jurisdiccion: al for al in actividad.alicuotas}
    for j in jurisdicciones:
        texto = (leer(j) or "").strip()
        if not texto:
            if j in actuales:
                db.session.delete(actuales[j])
            continue
        valor = _num(texto)
        if valor < 0 or valor > 100:
            raise ValueError(f"La alícuota de «{actividad.nombre}» en {j} tiene que estar entre 0 y 100.")
        if j in actuales:
            actuales[j].alicuota = valor
        else:
            db.session.add(CmAlicuota(actividad_id=actividad.id, jurisdiccion=j, alicuota=valor))


@cm_bp.route("/coeficientes", methods=["GET", "POST"])
@estudio_required
def coeficientes(cliente_id):
    cliente = _cliente_del_estudio_o_404(cliente_id)
    habilitadas = [j.jurisdiccion for j in _jurisdicciones_cliente(cliente)]
    if request.method == "POST":
        vigencia = request.form.get("vigencia_desde", "").strip()
        if not RE_PERIODO.match(vigencia):
            flash("La vigencia tiene que tener el formato AAAA-MM (por ejemplo 2026-04).")
            return redirect(url_for("cm.coeficientes", cliente_id=cliente.id))
        try:
            coefs = {}
            for j in habilitadas:
                texto = request.form.get("coef_" + j, "").strip()
                if texto:
                    coefs[j] = round(_num(texto), 4)
        except ValueError:
            flash("Algún coeficiente no es un número válido (ej. 0,3050).")
            return redirect(url_for("cm.coeficientes", cliente_id=cliente.id))
        error = validar_coeficientes({j: c for j, c in coefs.items() if c > 0})
        if error:
            flash(error)
            return redirect(url_for("cm.coeficientes", cliente_id=cliente.id))
        CmCoeficiente.query.filter_by(cliente_id=cliente.id, vigencia_desde=vigencia).delete()
        for j, c in coefs.items():
            if c > 0:
                db.session.add(CmCoeficiente(cliente_id=cliente.id, vigencia_desde=vigencia, jurisdiccion=j, coeficiente=c))
        db.session.commit()
        flash(f"Coeficientes vigentes desde {vigencia} guardados.")
        return redirect(url_for("cm.coeficientes", cliente_id=cliente.id))
    return render_template(
        "cm_coeficientes.html", cliente=cliente, habilitadas=habilitadas, conjuntos=_conjuntos_coeficientes(cliente),
        ultimo=(_conjuntos_coeficientes(cliente) or [(None, {})])[-1][1],
    )


@cm_bp.route("/coeficientes/<vigencia>/eliminar", methods=["POST"])
@estudio_required
def coeficientes_eliminar(cliente_id, vigencia):
    cliente = _cliente_del_estudio_o_404(cliente_id)
    _periodo_o_404(vigencia)
    CmCoeficiente.query.filter_by(cliente_id=cliente.id, vigencia_desde=vigencia).delete()
    db.session.commit()
    flash(f"Conjunto de coeficientes desde {vigencia} eliminado.")
    return redirect(url_for("cm.coeficientes", cliente_id=cliente.id))


@cm_bp.route("/<periodo>", methods=["GET", "POST"])
@estudio_required
def liquidacion(cliente_id, periodo):
    cliente = _cliente_del_estudio_o_404(cliente_id)
    _periodo_o_404(periodo)
    faltan = _faltantes_de_configuracion(cliente, periodo)
    if faltan:
        for m in faltan:
            flash(m)
        return redirect(url_for("cm.inicio", cliente_id=cliente.id))

    ctx = _contexto_periodo(cliente, periodo)
    if request.method == "POST":
        try:
            nuevos_ingresos = {a.id: _num(request.form.get(f"ing_{a.id}")) for a in ctx["actividades"]}
            nuevos_creditos = {
                j: {k: _num(request.form.get(f"cred_{j}_{k}")) for k in CAMPOS_CREDITO} for j in ctx["habilitadas"]
            }
        except ValueError:
            flash("Alguno de los importes no es un número válido.")
            return redirect(url_for("cm.liquidacion", cliente_id=cliente.id, periodo=periodo))
        # El saldo a favor anterior y las deducciones no pueden ser negativos; los ingresos sí (notas de crédito).
        if any(v < 0 for c in nuevos_creditos.values() for v in c.values()):
            flash("Las retenciones, percepciones, pagos y saldos a favor se cargan en positivo.")
            return redirect(url_for("cm.liquidacion", cliente_id=cliente.id, periodo=periodo))
        CmIngreso.query.filter_by(cliente_id=cliente.id, periodo=periodo).delete()
        for aid, imp in nuevos_ingresos.items():
            db.session.add(CmIngreso(cliente_id=cliente.id, periodo=periodo, actividad_id=aid, importe=imp))
        CmCredito.query.filter_by(cliente_id=cliente.id, periodo=periodo).delete()
        for j, c in nuevos_creditos.items():
            db.session.add(CmCredito(cliente_id=cliente.id, periodo=periodo, jurisdiccion=j, **c))
        db.session.commit()
        flash("Datos del período guardados.")
        return redirect(url_for("cm.liquidacion", cliente_id=cliente.id, periodo=periodo))

    resultado = _calcular(cliente, periodo, ctx)
    asignado = round(sum(ctx["ingresos"].values()), 2)
    return render_template(
        "cm_liquidacion.html", cliente=cliente, periodo=periodo, ctx=ctx, resultado=resultado,
        diferencia_ventas=round(ctx["ventas"] - asignado, 2),
    )


@cm_bp.route("/<periodo>/pdf")
@estudio_required
def descargar_pdf(cliente_id, periodo):
    cliente = _cliente_del_estudio_o_404(cliente_id)
    _periodo_o_404(periodo)
    if _faltantes_de_configuracion(cliente, periodo):
        return redirect(url_for("cm.liquidacion", cliente_id=cliente.id, periodo=periodo))
    from cm_pdf import generar_pdf_cm
    ctx = _contexto_periodo(cliente, periodo)
    resultado = _calcular(cliente, periodo, ctx)
    path = os.path.join(TMP_DIR, f"cm_{current_user.id}_{uuid.uuid4().hex}.pdf")
    try:
        generar_pdf_cm(resultado, cliente, path, estudio=current_user.nombre)
        with open(path, "rb") as fh:
            datos = io.BytesIO(fh.read())
    finally:
        if os.path.exists(path):
            os.remove(path)
    return send_file(datos, mimetype="application/pdf", as_attachment=True, download_name=f"CM03_{periodo}_{cliente.cuit}.pdf")


def _filas_percepciones(cliente, periodo):
    """Una fila por percepción de IIBB de las facturas de compra del período (para controlar o importar)."""
    habilitadas = {j.jurisdiccion for j in cliente.cm_jurisdicciones}
    filas = []
    for c in Comprobante.query.filter_by(cliente_id=cliente.id, tipo="compra", periodo=periodo).order_by(Comprobante.fecha, Comprobante.id):
        signo = -1 if _es_nota_de_credito(c) else 1
        pares = [(p.jurisdiccion, float(p.importe)) for p in c.percepciones_iibb]
        if not pares and float(c.percepcion_iibb or 0):
            pares = [(c.jurisdiccion_iibb, float(c.percepcion_iibb))]
        letra = (c.tipo_comprobante or "").strip().split(" ")[-1][:1].upper()
        for nombre, importe in pares:
            jur = canonizar_jurisdiccion(nombre)
            filas.append({
                "jurisdiccion": jur or "", "texto_original": nombre or "", "en_convenio": jur in habilitadas,
                "cuit_agente": c.cuit, "razon_social": c.razon_social, "fecha": c.fecha,
                "comprobante": c.tipo_comprobante, "letra": letra, "punto_venta": c.punto_venta,
                "numero": c.numero_comprobante, "importe": round(signo * importe, 2),
            })
    return filas


@cm_bp.route("/<periodo>/percepciones.txt")
@estudio_required
def descargar_percepciones_txt(cliente_id, periodo):
    """TXT para importar las percepciones de IIBB en SIFERE web (Deducciones → Percepciones → Importar)."""
    cliente = _cliente_del_estudio_o_404(cliente_id)
    _periodo_o_404(periodo)
    from sifere_txt import generar_txt
    texto, omitidas = generar_txt(_filas_percepciones(cliente, periodo))
    for desc, motivo in omitidas:
        flash(f"No incluida en el TXT: {desc} — {motivo}.")
    if not texto:
        flash(f"No hay percepciones para exportar en {periodo}.")
        return redirect(url_for("cm.liquidacion", cliente_id=cliente.id, periodo=periodo))
    if omitidas:
        flash("Descargaste el TXT con las que sí se pudieron armar; cargá a mano las que quedaron afuera.")
    return send_file(
        io.BytesIO(texto.encode("ascii", errors="replace")), mimetype="text/plain", as_attachment=True,
        download_name=f"Percepciones_SIFERE_{periodo}_{cliente.cuit}.txt",
    )


@cm_bp.route("/<periodo>/percepciones.xlsx")
@estudio_required
def descargar_percepciones(cliente_id, periodo):
    """Detalle de percepciones de IIBB sufridas (de las facturas de compra escaneadas) en Excel."""
    cliente = _cliente_del_estudio_o_404(cliente_id)
    _periodo_o_404(periodo)
    from openpyxl import Workbook
    from openpyxl.styles import Font
    filas = _filas_percepciones(cliente, periodo)
    if not filas:
        flash(f"No hay percepciones de IIBB en las facturas de compra de {periodo}.")
        return redirect(url_for("cm.liquidacion", cliente_id=cliente.id, periodo=periodo))
    wb = Workbook()
    ws = wb.active
    ws.title = "Percepciones IIBB"
    ws.append(["Jurisdicción", "En el Convenio del cliente", "CUIT del agente (proveedor)", "Razón social", "Fecha",
               "Comprobante", "Letra", "Punto de venta", "Número", "Importe percibido"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for f in filas:
        ws.append([f["jurisdiccion"] or f"SIN IDENTIFICAR ({f['texto_original'] or 'vacío'})", "Sí" if f["en_convenio"] else "No",
                   f["cuit_agente"], f["razon_social"], f["fecha"], f["comprobante"], f["letra"],
                   f"{f['punto_venta']:04d}", f"{f['numero']:08d}", f["importe"]])
    n = len(filas) + 1
    ws.append([]); ws.append(["TOTAL", "", "", "", "", "", "", "", "", f"=SUM(J2:J{n})"])
    ws.cell(n + 2, 1).font = Font(bold=True)
    for row in ws.iter_rows(min_row=2, max_row=n + 2, min_col=5, max_col=5):
        row[0].number_format = "dd/mm/yyyy"
    for row in ws.iter_rows(min_row=2, max_row=n + 2, min_col=10, max_col=10):
        row[0].number_format = "#,##0.00"
    for col, w in zip("ABCDEFGHIJ", (24, 14, 22, 36, 12, 22, 7, 12, 12, 16)):
        ws.column_dimensions[col].width = w
    datos = io.BytesIO()
    wb.save(datos)
    datos.seek(0)
    return send_file(
        datos, as_attachment=True, download_name=f"Percepciones_IIBB_{periodo}_{cliente.cuit}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@cm_bp.route("/<periodo>/xlsx")
@estudio_required
def descargar_xlsx(cliente_id, periodo):
    cliente = _cliente_del_estudio_o_404(cliente_id)
    _periodo_o_404(periodo)
    if _faltantes_de_configuracion(cliente, periodo):
        return redirect(url_for("cm.liquidacion", cliente_id=cliente.id, periodo=periodo))
    from cm_pdf import generar_xlsx_cm
    resultado = _calcular(cliente, periodo)
    datos = io.BytesIO()
    generar_xlsx_cm(resultado, cliente, datos)
    datos.seek(0)
    return send_file(
        datos, as_attachment=True, download_name=f"CM03_{periodo}_{cliente.cuit}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
