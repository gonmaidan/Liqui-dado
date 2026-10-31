"""
cm_pdf.py
-----------
PDF y Excel de trabajo del CM03 del período: ingresos por actividad, distribución por
jurisdicción con el detalle del cálculo, deducciones y saldo a pagar / a favor.
Es un papel de trabajo; la declaración se presenta en SIFERE web.
"""

from datetime import date

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

AZUL = colors.HexColor("#1F3A5F")
GRIS = colors.HexColor("#F2F4F8")
BORDE = colors.HexColor("#C9D0DD")
SUAVE = colors.HexColor("#5A6483")


def _m(v: float) -> str:
    t = f"{abs(v):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return ("-" if v < 0 else "") + t


def _pie(canvas, doc):
    canvas.saveState()
    canvas.setFont("Helvetica", 7.5)
    canvas.setFillColor(SUAVE)
    ancho, _ = landscape(A4)
    canvas.drawString(1.5 * cm, 1 * cm, "Papel de trabajo generado por app-liquidar. Se presenta en SIFERE web; controlar contra la declaración presentada.")
    canvas.drawRightString(ancho - 1.5 * cm, 1 * cm, f"Página {doc.page}")
    canvas.restoreState()


def _estilo_tabla(extra=()):
    base = [
        ("BACKGROUND", (0, 0), (-1, 0), AZUL), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("ALIGN", (1, 0), (-1, -1), "RIGHT"), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, GRIS]),
        ("GRID", (0, 0), (-1, -1), 0.4, BORDE), ("TOPPADDING", (0, 0), (-1, -1), 3.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 3.5),
    ]
    return TableStyle(base + list(extra))


def generar_pdf_cm(r, cliente, path, estudio="Estudio Contable"):
    doc = SimpleDocTemplate(
        path, pagesize=landscape(A4), topMargin=1.4 * cm, bottomMargin=1.6 * cm, leftMargin=1.5 * cm, rightMargin=1.5 * cm,
        title=f"CM03 {r.periodo} - {cliente.razon_social}",
    )
    est = getSampleStyleSheet()
    titulo = ParagraphStyle("t", parent=est["Title"], fontName="Helvetica-Bold", fontSize=17, textColor=AZUL, alignment=0, spaceAfter=2)
    sub = ParagraphStyle("s", parent=est["Normal"], fontSize=9.5, textColor=SUAVE, leading=13, spaceAfter=8)
    h2 = ParagraphStyle("h2", parent=est["Heading2"], fontName="Helvetica-Bold", fontSize=11, textColor=AZUL, spaceBefore=10, spaceAfter=5)
    chico = ParagraphStyle("c", parent=est["Normal"], fontSize=8, textColor=SUAVE, leading=11)

    anio, mes = r.periodo.split("-")
    story = [
        Paragraph(f"Convenio Multilateral — CM03 {mes}/{anio}", titulo),
        Paragraph(f"{estudio} · {cliente.razon_social} · CUIT {cliente.cuit} · Coeficientes unificados vigentes desde {r.vigencia_coef}", sub),
        Paragraph("Ingresos computables del mes", h2),
    ]
    filas = [["Actividad", "Ingreso del mes ($)"]] + [[a, _m(v)] for a, v in r.ingresos.items()] + [["TOTAL", _m(r.total_ingresos)]]
    t = Table(filas, colWidths=[12 * cm, 5 * cm], hAlign="LEFT")
    t.setStyle(_estilo_tabla([("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"), ("ALIGN", (0, 0), (0, -1), "LEFT")]))
    story += [t, Paragraph("Distribución por jurisdicción", h2)]

    cab = ["Jurisdicción", "Coef.", "Base\nimponible", "Imp. por\nalícuota", "Mínimo", "Imp.\ndeterminado",
           "Retenc.", "SIRCREB", "Percep.", "Pagos a\ncuenta", "Saldo a favor\nanterior", "Saldo\na pagar / (a favor)"]
    filas = [cab]
    for f in r.filas:
        filas.append([f.jurisdiccion, f"{f.coeficiente:.4f}".replace(".", ","), _m(f.base), _m(f.impuesto), _m(f.minimo) if f.minimo else "-",
                      _m(f.impuesto_determinado), _m(f.retenciones), _m(f.sircreb), _m(f.percepciones), _m(f.pagos),
                      _m(f.saldo_anterior), _m(f.saldo) if f.saldo >= 0 else f"({_m(-f.saldo)})"])
    tt = r.totales
    filas.append(["TOTALES", "", _m(tt["base"]), "", "", _m(tt["impuesto_determinado"]), "", "", "", "", "", ""])
    tabla = Table(filas, colWidths=[3.2 * cm, 1.4 * cm] + [2.2 * cm] * 10, repeatRows=1)
    tabla.setStyle(_estilo_tabla([("ALIGN", (0, 0), (0, -1), "LEFT"), ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
                                  ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#E1E7F2")), ("FONTSIZE", (0, 0), (-1, 0), 7)]))
    story.append(tabla)

    story.append(Paragraph("Detalle por actividad y alícuota", h2))
    det = [["Jurisdicción", "Actividad", "Base ($)", "Alícuota", "Impuesto ($)"]]
    for f in r.filas:
        for act, base, alic, imp in f.detalle:
            det.append([f.jurisdiccion, act, _m(base), f"{alic:g}%".replace(".", ","), _m(imp)])
    td = Table(det, colWidths=[4 * cm, 9 * cm, 3.5 * cm, 2.5 * cm, 3.5 * cm], repeatRows=1)
    td.setStyle(_estilo_tabla([("ALIGN", (0, 0), (1, -1), "LEFT")]))
    story.append(td)

    story.append(Paragraph("Resumen", h2))
    res = [["Concepto", "Importe ($)"],
           ["Impuesto determinado total", _m(tt["impuesto_determinado"])],
           ["Deducciones (retenciones, percepciones, pagos y saldos a favor)", _m(tt["deducciones"])],
           ["TOTAL A PAGAR (suma de las jurisdicciones con saldo)", _m(tt["a_pagar"])],
           ["Saldos a favor que quedan por jurisdicción", _m(tt["a_favor"])]]
    tr = Table(res, colWidths=[13 * cm, 4.5 * cm], hAlign="LEFT")
    tr.setStyle(_estilo_tabla([("ALIGN", (0, 0), (0, -1), "LEFT"), ("FONTNAME", (0, 3), (-1, 3), "Helvetica-Bold"),
                               ("BACKGROUND", (0, 3), (-1, 3), colors.HexColor("#E1E7F2"))]))
    story.append(tr)

    notas = [f"Los saldos a favor quedan en su jurisdicción y no se compensan con saldos a pagar de otra."] + list(r.avisos)
    for f in r.filas:
        notas += [f"{f.jurisdiccion}: {a}" for a in f.avisos]
    story.append(Spacer(1, 8))
    for n in notas:
        story.append(Paragraph("• " + n, chico))
    story.append(Paragraph(f"Generado el {date.today():%d/%m/%Y}. Alícuotas, mínimos y coeficientes según lo cargado por el estudio.", chico))
    doc.build(story, onFirstPage=_pie, onLaterPages=_pie)


def generar_xlsx_cm(r, cliente, salida):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "CM03"
    ws["A1"] = f"CM03 {r.periodo} — {cliente.razon_social} (CUIT {cliente.cuit})"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = f"Coeficientes unificados vigentes desde {r.vigencia_coef}"

    fila = 4
    ws.cell(fila, 1, "Actividad").font = Font(bold=True)
    ws.cell(fila, 2, "Ingreso del mes").font = Font(bold=True)
    for a, v in r.ingresos.items():
        fila += 1
        ws.cell(fila, 1, a)
        ws.cell(fila, 2, v)
    fila += 1
    ws.cell(fila, 1, "TOTAL").font = Font(bold=True)
    ws.cell(fila, 2, f"=SUM(B5:B{fila - 1})").font = Font(bold=True)

    fila += 2
    cab = ["Jurisdicción", "Coeficiente", "Base imponible", "Impuesto por alícuota", "Mínimo", "Impuesto determinado",
           "Retenciones", "SIRCREB", "Percepciones", "Pagos a cuenta", "Saldo a favor anterior", "Saldo (+ a pagar / - a favor)"]
    cabecera = fila
    for i, t in enumerate(cab, 1):
        c = ws.cell(fila, i, t)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3A5F")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for f in r.filas:
        fila += 1
        ws.cell(fila, 1, f.jurisdiccion)
        ws.cell(fila, 2, f.coeficiente)
        ws.cell(fila, 3, f.base)
        ws.cell(fila, 4, f.impuesto)
        ws.cell(fila, 5, f.minimo)
        ws.cell(fila, 6, f"=MAX(D{fila},E{fila})")
        for col, v in zip(range(7, 12), (f.retenciones, f.sircreb, f.percepciones, f.pagos, f.saldo_anterior)):
            ws.cell(fila, col, v)
        ws.cell(fila, 12, f"=F{fila}-SUM(G{fila}:K{fila})")
    fila += 1
    ws.cell(fila, 1, "TOTALES").font = Font(bold=True)
    for col in (3, 6, 12):
        letra = ws.cell(fila, col).column_letter
        ws.cell(fila, col, f"=SUM({letra}{cabecera + 1}:{letra}{fila - 1})").font = Font(bold=True)
    for row in ws.iter_rows(min_row=5, max_row=fila, min_col=2, max_col=12):
        for c in row:
            c.number_format = "0.0000" if c.column == 2 and c.row > cabecera else "#,##0.00"
    ws.column_dimensions["A"].width = 26
    for col in "BCDEFGHIJKL":
        ws.column_dimensions[col].width = 18

    ws2 = wb.create_sheet("Detalle")
    ws2.append(["Jurisdicción", "Actividad", "Base", "Alícuota %", "Impuesto"])
    for f in r.filas:
        for act, base, alic, imp in f.detalle:
            ws2.append([f.jurisdiccion, act, base, alic, imp])
    for c in ws2[1]:
        c.font = Font(bold=True)
    ws2.column_dimensions["A"].width = 22
    ws2.column_dimensions["B"].width = 40
    for col in "CDE":
        ws2.column_dimensions[col].width = 16
    wb.save(salida)
